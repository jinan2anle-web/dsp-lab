# -*- coding: utf-8 -*-
"""打印周报模板结构，确认要填充的形状索引与单元格字体"""
import sys
from pptx import Presentation
from pptx.util import Emu

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

P = r'C:\Users\jinan2an\OneDrive\Desktop\dsh_work_area\task1_w1\周报模板.pptx'
prs = Presentation(P)
print('slide size:', prs.slide_width, prs.slide_height,
      '=', Emu(prs.slide_width).inches, 'x', Emu(prs.slide_height).inches, 'in')
for i, sh in enumerate(prs.slides[0].shapes):
    print('---', i, sh.shape_type, sh.name, 'has_tf=', sh.has_text_frame,
          'has_table=', sh.has_table)
    print('    pos', sh.left, sh.top, 'size', sh.width, sh.height)
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                print('    TEXT %r size=%s bold=%s font=%s' % (
                    r.text, r.font.size, r.font.bold, r.font.name))
    if sh.has_table:
        t = sh.table
        print('    rows=%d cols=%d' % (len(t.rows), len(t.columns)))
        for ci, c in enumerate(t.columns):
            print('    col%d width=%d' % (ci, c.width))
        for ri, row in enumerate(t.rows):
            print('    row%d height=%d' % (ri, row.height))
            for ci, cell in enumerate(row.cells):
                txt = cell.text
                r0 = None
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        r0 = r
                        break
                    break
                print('      [%d][%d] %r size=%s bold=%s font=%s' % (
                    ri, ci, txt[:40],
                    r0.font.size if r0 else None,
                    r0.font.bold if r0 else None,
                    r0.font.name if r0 else None))
