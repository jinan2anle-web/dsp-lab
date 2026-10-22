# -*- coding: utf-8 -*-
"""导出周报 PPTX 的全部文字，用于内容终检"""
import sys
from pptx import Presentation

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

P = r'C:\Users\jinan2an\OneDrive\Desktop\dsh_work_area\task1_w1\第一周周报.pptx'
prs = Presentation(P)
for si, slide in enumerate(prs.slides):
    print('=' * 70)
    print('slide %d' % (si + 1))
    print('=' * 70)
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            print('[TEXT] %s' % sh.text_frame.text)
        if sh.has_table:
            t = sh.table
            for ri, row in enumerate(t.rows):
                for ci, cell in enumerate(row.cells):
                    if not cell.text.strip():
                        continue
                    print('[R%dC%d] %s' % (ri, ci, cell.text.replace('\n', ' ⏎ ')))
