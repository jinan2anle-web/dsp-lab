# -*- coding: utf-8 -*-
"""估算周报表格各单元格文本是否超出可用高度（中文按 1 字宽、西文按 0.55 字宽）"""
import sys
from pptx import Presentation
from pptx.util import Emu, Pt

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

P = r'C:\Users\jinan2an\OneDrive\Desktop\dsh_work_area\task1_w1\第一周周报.pptx'
SLIDE_H = 7.5 * 914400
TABLE_TOP = 739775
TABLE_H = 5784931
HEADER_H = 736600

prs = Presentation(P)
slide = prs.slides[0]
for sh in slide.shapes:
    if not sh.has_table:
        continue
    t = sh.table
    print('table at top=%.2f in, height=%.2f in, bottom=%.2f in (slide 7.5 in)'
          % (Emu(sh.top).inches, Emu(sh.height).inches,
             (sh.top + sh.height) / 914400.0))
    body_h = TABLE_H - HEADER_H
    print('body row available height = %.2f in' % (body_h / 914400.0))
    for ci, col in enumerate(t.columns):
        w_in = Emu(col.width).inches
        cell = t.cell(1, ci)
        total_lines = 0.0
        chars_avail = (w_in - 0.14) * 72 / 11.0 * 1.0   # 内边距后按 11pt 全角估算
        for p in cell.text_frame.paragraphs:
            txt = ''.join(r.text for r in p.runs)
            if not txt:
                continue
            for seg in txt.split('\n'):
                width = sum(1.0 if ord(c) > 0x2000 else 0.55 for c in seg)
                total_lines += max(1.0, -(-width // max(chars_avail, 1)))
        need_in = total_lines * 11 * 1.2 / 72.0
        print('  col%d w=%.2f in  行数≈%.1f  需高≈%.2f in  %s'
              % (ci, w_in, total_lines, need_in,
                 'OK' if need_in <= body_h / 914400.0 else '*** 可能溢出 ***'))
